import csv
import json
import defusedxml.ElementTree as ET
import openpyxl

def parse_csv(file_path: str) -> list[dict]:
    with open(file_path, mode='r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        return [row for row in reader]

def parse_json(file_path: str) -> list[dict]:
    with open(file_path, mode='r', encoding='utf-8') as f:
        data = json.load(f)
        if isinstance(data, dict):
            return [data]
        elif isinstance(data, list):
            return data
        else:
            raise ValueError("JSON root must be object or array")

def parse_xml(file_path: str) -> list[dict]:
    tree = ET.parse(file_path)
    root = tree.getroot()
    records = []
    for child in root:
        record = {}
        for elem in child:
            record[elem.tag] = elem.text
        records.append(record)
    return records

def parse_txt(file_path: str) -> list[dict]:
    # Roadmap requirement says TXT must follow predefined structure.
    # We will assume a simple tab-separated or structured lines.
    # For now, if undefined, we stage as raw lines.
    records = []
    with open(file_path, mode='r', encoding='utf-8') as f:
        for line in f:
            stripped = line.strip()
            if stripped:
                records.append({"raw_line": stripped})
    return records

def _normalize_cell_value(val):
    if val is None:
        return None
    import datetime
    if isinstance(val, datetime.datetime):
        if val.time() == datetime.time(0, 0, 0):
            return val.date().isoformat()
        return val.isoformat()
    if isinstance(val, datetime.date):
        return val.isoformat()
    return val

def parse_xlsx(file_path: str) -> list[dict]:
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        all_records = []
        for sheetname in wb.sheetnames:
            sheet = wb[sheetname]
            rows_iter = sheet.iter_rows(values_only=True)
            
            header_row = None
            header_row_idx = 0
            for r_idx, row in enumerate(rows_iter, start=1):
                if row and any(cell is not None and str(cell).strip() != "" for cell in row):
                    header_row = row
                    header_row_idx = r_idx
                    break
                    
            if not header_row:
                continue
                
            last_header_col = 0
            for idx, h in enumerate(header_row):
                if h is not None and str(h).strip() != "":
                    last_header_col = idx + 1
                    
            if last_header_col == 0:
                continue
                
            raw_headers = list(header_row[:last_header_col])
            
            seen_headers = set()
            clean_headers = []
            for h in raw_headers:
                if h is None or str(h).strip() == "":
                    clean_h = f"Unnamed_{len(clean_headers)}"
                else:
                    clean_h = str(h)
                    if clean_h in seen_headers:
                        raise ValueError(f"Duplicate header '{clean_h}' found in worksheet '{sheetname}'")
                    seen_headers.add(clean_h)
                clean_headers.append(clean_h)
                
            current_row_idx = header_row_idx
            for row in rows_iter:
                current_row_idx += 1
                if not row:
                    continue
                if not any(c is not None and (not isinstance(c, str) or c.strip() != "") for c in row):
                    continue
                    
                row_slice = row[:len(clean_headers)]
                if len(row_slice) < len(clean_headers):
                    row_slice = list(row_slice) + [None] * (len(clean_headers) - len(row_slice))
                    
                row_data = {}
                for h_name, cell_val in zip(clean_headers, row_slice):
                    row_data[h_name] = _normalize_cell_value(cell_val)
                    
                all_records.append({
                    "sheet_name": sheetname,
                    "row_number": current_row_idx,
                    "data": row_data
                })
                
        return all_records
    finally:
        wb.close()

def parse_file(file_path: str, file_type: str) -> list[dict]:
    parsers = {
        "csv": parse_csv,
        "json": parse_json,
        "xml": parse_xml,
        "txt": parse_txt,
        "xlsx": parse_xlsx
    }
    
    parser = parsers.get(file_type.lower())
    if not parser:
        raise ValueError(f"Unsupported file type: {file_type}")
        
    return parser(file_path)
