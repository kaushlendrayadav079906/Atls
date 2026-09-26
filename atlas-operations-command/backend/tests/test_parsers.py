import pytest
import os
import json
import datetime
from openpyxl import Workbook

from app.services.parsers import (
    parse_csv,
    parse_json,
    parse_xml,
    parse_txt,
    parse_xlsx,
    parse_file
)

def test_parse_csv(tmp_path):
    p = tmp_path / "test.csv"
    p.write_text("name,amount\nFactory1,1000\nFactory2,2000")
    records = parse_csv(str(p))
    assert len(records) == 2
    assert records[0]["name"] == "Factory1"
    assert records[0]["amount"] == "1000"

def test_parse_json(tmp_path):
    p = tmp_path / "test.json"
    p.write_text(json.dumps([{"name": "Factory1"}, {"name": "Factory2"}]))
    records = parse_json(str(p))
    assert len(records) == 2
    assert records[1]["name"] == "Factory2"

def test_parse_xml(tmp_path):
    p = tmp_path / "test.xml"
    p.write_text("""<?xml version="1.0"?>
    <root>
        <record><name>Factory1</name><amount>1000</amount></record>
        <record><name>Factory2</name><amount>2000</amount></record>
    </root>""")
    records = parse_xml(str(p))
    assert len(records) == 2
    assert records[0]["name"] == "Factory1"
    assert records[0]["amount"] == "1000"

def test_parse_txt(tmp_path):
    p = tmp_path / "test.txt"
    p.write_text("Line 1\nLine 2\n\nLine 3")
    records = parse_txt(str(p))
    assert len(records) == 3
    assert records[0]["raw_line"] == "Line 1"

def test_parse_file_factory(tmp_path):
    p = tmp_path / "test.json"
    p.write_text(json.dumps([{"test": 1}]))
    res = parse_file(str(p), "json")
    assert len(res) == 1
    assert res[0]["test"] == 1

# =========================================================================
# Generic Multi-Sheet XLSX Parser Unit Tests (Step 3)
# =========================================================================

def test_parse_xlsx_single_sheet(tmp_path):
    """A. Verify single-sheet parsing preserves sheet_name, row_number, and data dict."""
    p = tmp_path / "single_sheet.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Factories"
    ws.append(["name", "amount"])
    ws.append(["Factory1", 1000])
    ws.append(["Factory2", 2000])
    wb.save(str(p))
    
    records = parse_xlsx(str(p))
    assert len(records) == 2
    assert records[0]["sheet_name"] == "Factories"
    assert records[0]["row_number"] == 2
    assert records[0]["data"]["name"] == "Factory1"
    assert records[0]["data"]["amount"] == 1000

    assert records[1]["sheet_name"] == "Factories"
    assert records[1]["row_number"] == 3
    assert records[1]["data"]["name"] == "Factory2"
    assert records[1]["data"]["amount"] == 2000

def test_parse_xlsx_multi_sheet(tmp_path):
    """B. Verify multi-sheet workbook parsing processes all sheets."""
    p = tmp_path / "multi_sheet.xlsx"
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["Code", "Value"])
    ws1.append(["S1_R1", 10])
    
    ws2 = wb.create_sheet(title="Sheet2")
    ws2.append(["Code", "Value"])
    ws2.append(["S2_R1", 20])
    ws2.append(["S2_R2", 30])
    
    wb.save(str(p))
    
    records = parse_xlsx(str(p))
    assert len(records) == 3
    assert records[0]["sheet_name"] == "Sheet1"
    assert records[0]["data"]["Code"] == "S1_R1"
    assert records[1]["sheet_name"] == "Sheet2"
    assert records[1]["data"]["Code"] == "S2_R1"
    assert records[2]["sheet_name"] == "Sheet2"
    assert records[2]["data"]["Code"] == "S2_R2"

def test_parse_xlsx_sheet_order_preserved(tmp_path):
    """C & D. Verify worksheet order and names are strictly preserved."""
    p = tmp_path / "sheet_order.xlsx"
    wb = Workbook()
    ws_alpha = wb.active
    ws_alpha.title = "Alpha"
    ws_alpha.append(["col"])
    ws_alpha.append(["val_alpha"])
    
    ws_beta = wb.create_sheet(title="Beta")
    ws_beta.append(["col"])
    ws_beta.append(["val_beta"])
    
    ws_gamma = wb.create_sheet(title="Gamma")
    ws_gamma.append(["col"])
    ws_gamma.append(["val_gamma"])
    
    wb.save(str(p))
    
    records = parse_xlsx(str(p))
    assert [r["sheet_name"] for r in records] == ["Alpha", "Beta", "Gamma"]

def test_parse_xlsx_row_order_and_row_number(tmp_path):
    """E & G. Verify row ordering and 1-based row numbers are accurately tracked."""
    p = tmp_path / "row_order.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Ordering"
    # Row 1 is header
    ws.append(["Seq", "Desc"])
    # Row 2, 3, 4 are data
    ws.append([1, "First"])
    ws.append([2, "Second"])
    ws.append([3, "Third"])
    wb.save(str(p))
    
    records = parse_xlsx(str(p))
    assert len(records) == 3
    assert [r["row_number"] for r in records] == [2, 3, 4]
    assert [r["data"]["Seq"] for r in records] == [1, 2, 3]
    assert [r["data"]["Desc"] for r in records] == ["First", "Second", "Third"]

def test_parse_xlsx_exact_headers_preserved(tmp_path):
    """F. Verify exact preservation of complex headers (spaces, punctuation, casing, parentheses, symbols)."""
    p = tmp_path / "headers.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Headers"
    exact_headers = ["Invoice No.", "CGST (Rs.)", "Margin %", "Bill to", "Item Code/SAC", "Freight (Tax Amt.)"]
    ws.append(exact_headers)
    ws.append(["INV-1001", 180.50, 24.5, "123 Main St", "HSN3926", 45.00])
    wb.save(str(p))
    
    records = parse_xlsx(str(p))
    assert len(records) == 1
    data = records[0]["data"]
    for header in exact_headers:
        assert header in data
    assert data["Invoice No."] == "INV-1001"
    assert data["CGST (Rs.)"] == 180.50
    assert data["Margin %"] == 24.5
    assert data["Bill to"] == "123 Main St"
    assert data["Item Code/SAC"] == "HSN3926"
    assert data["Freight (Tax Amt.)"] == 45.00

def test_parse_xlsx_datetime_and_date_handling(tmp_path):
    """H & I. Verify datetime and date handling normalized to ISO strings."""
    p = tmp_path / "dates.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Dates"
    ws.append(["Timestamp", "CalendarDate", "PlainString"])
    
    dt_val = datetime.datetime(2026, 3, 15, 14, 30, 45)
    d_val = datetime.date(2026, 3, 15)
    ws.append([dt_val, d_val, "2026-03-15"])
    wb.save(str(p))
    
    records = parse_xlsx(str(p))
    assert len(records) == 1
    data = records[0]["data"]
    assert data["Timestamp"] == "2026-03-15T14:30:45"
    assert data["CalendarDate"] == "2026-03-15"
    assert data["PlainString"] == "2026-03-15"

def test_parse_xlsx_empty_cells_and_nulls(tmp_path):
    """J. Verify empty Excel cells evaluate to None without conversion to empty string or 0."""
    p = tmp_path / "empty_cells.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sparse"
    ws.append(["ColA", "ColB", "ColC"])
    ws.append(["ValA", None, "ValC"])
    ws.append([None, None, "OnlyC"])
    wb.save(str(p))
    
    records = parse_xlsx(str(p))
    assert len(records) == 2
    assert records[0]["data"]["ColA"] == "ValA"
    assert records[0]["data"]["ColB"] is None
    assert records[0]["data"]["ColC"] == "ValC"
    
    assert records[1]["data"]["ColA"] is None
    assert records[1]["data"]["ColB"] is None
    assert records[1]["data"]["ColC"] == "OnlyC"

def test_parse_xlsx_numeric_types(tmp_path):
    """K & L. Verify integers, floats, and numeric types are preserved accurately."""
    p = tmp_path / "numerics.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Financial"
    ws.append(["Quantity", "UnitPrice", "GrossAmount"])
    ws.append([150, 1250.75, 187612.50])
    wb.save(str(p))
    
    records = parse_xlsx(str(p))
    assert len(records) == 1
    data = records[0]["data"]
    assert isinstance(data["Quantity"], int)
    assert data["Quantity"] == 150
    assert isinstance(data["UnitPrice"], float)
    assert data["UnitPrice"] == 1250.75
    assert data["GrossAmount"] == 187612.50

def test_parse_xlsx_empty_worksheet_handling(tmp_path):
    """M. Verify completely empty sheets or sheets with only headers do not cause crashes."""
    p = tmp_path / "empty_sheets.xlsx"
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "EmptySheet1"
    # No rows at all
    
    ws2 = wb.create_sheet(title="HeaderOnly")
    ws2.append(["HeaderA", "HeaderB"])
    # No data rows
    
    ws3 = wb.create_sheet(title="ValidSheet")
    ws3.append(["ID", "Name"])
    ws3.append([101, "Item1"])
    
    wb.save(str(p))
    
    records = parse_xlsx(str(p))
    assert len(records) == 1
    assert records[0]["sheet_name"] == "ValidSheet"
    assert records[0]["data"]["ID"] == 101

def test_parse_xlsx_duplicate_header_rejection(tmp_path):
    """N. Verify duplicate headers raise explicit ValueError."""
    p = tmp_path / "dup_headers.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "BadHeaders"
    ws.append(["Invoice No", "Amount", "Amount"]) # Duplicate 'Amount'
    ws.append(["INV-1", 100, 200])
    wb.save(str(p))
    
    with pytest.raises(ValueError, match="Duplicate header 'Amount' found in worksheet 'BadHeaders'"):
        parse_xlsx(str(p))

def test_parse_xlsx_astroform_synthetic_multi_sheet(tmp_path):
    """8. Section 8: Synthetic Astroform 3-sheet workbook parsing verification."""
    p = tmp_path / "synthetic_astroform.xlsx"
    wb = Workbook()
    
    # Sheet 1: Margin Report
    ws_margin = wb.active
    ws_margin.title = "Margin Report"
    ws_margin.append([
        "Invoice No.", "Invoice Date", "Customer Code", "Customer Name",
        "Item No.", "Item Description", "Sales Quantity", "Basic Amount",
        "Total RM Cost", "Freight", "Discount", "Margin", "Margin %", "Loc"
    ])
    ws_margin.append([
        26270441, datetime.date(2026, 3, 10), "C0070", "Customer Alpha",
        "ITM-01", "Precision Part A", 500, 150000.0,
        90000.0, 5000.0, 0.0, 55000.0, 36.67, "PUNE"
    ])
    
    # Sheet 2: Sale Register Customer Wise
    ws_sales = wb.create_sheet(title="Sale Register Customer Wise")
    ws_sales.append([
        "Invoice No", "Invoice Date", "Customer Code", "Customer Name",
        "Item Code", "Item Name", "Basic Amount", "CGST Rate", "CGSTAmt",
        "SGST Rate", "SGSTAmt", "IGST Rate", "IGSTAmt", "Total Invoice Amount", "Loc"
    ])
    ws_sales.append([
        "INV-26270441", datetime.date(2026, 3, 10), "C0070", "Customer Alpha",
        "ITM-01", "Precision Part A", 150000.0, 9.0, 13500.0,
        9.0, 13500.0, 0.0, 0.0, 177000.0, "PUNE"
    ])
    
    # Sheet 3: Purchase Register Vendor Wise
    ws_purchases = wb.create_sheet(title="Purchase Register Vendor Wise")
    ws_purchases.append([
        "Bill No", "Bill Date", "Vendor Name", "GST Regn No of BP",
        "Basic Amount", "CGST (Rs.)", "SGST (Rs.)", "IGST (Rs.)",
        "Gross Amount", "TDS (Rs.)", "loc"
    ])
    ws_purchases.append([
        "BILL-9901", datetime.date(2026, 3, 5), "Vendor Beta Steel", "27AAACB1234F1Z9",
        80000.0, 7200.0, 7200.0, 0.0,
        94400.0, 800.0, "PUNE"
    ])
    
    wb.save(str(p))
    
    records = parse_xlsx(str(p))
    assert len(records) == 3
    
    # Verify Sheet 1: Margin Report
    assert records[0]["sheet_name"] == "Margin Report"
    assert records[0]["row_number"] == 2
    assert records[0]["data"]["Invoice No."] == 26270441
    assert records[0]["data"]["Customer Code"] == "C0070"
    assert records[0]["data"]["Basic Amount"] == 150000.0
    assert records[0]["data"]["Margin %"] == 36.67
    
    # Verify Sheet 2: Sale Register Customer Wise
    assert records[1]["sheet_name"] == "Sale Register Customer Wise"
    assert records[1]["row_number"] == 2
    assert records[1]["data"]["Invoice No"] == "INV-26270441"
    assert records[1]["data"]["CGSTAmt"] == 13500.0
    assert records[1]["data"]["Total Invoice Amount"] == 177000.0
    
    # Verify Sheet 3: Purchase Register Vendor Wise
    assert records[2]["sheet_name"] == "Purchase Register Vendor Wise"
    assert records[2]["row_number"] == 2
    assert records[2]["data"]["Bill No"] == "BILL-9901"
    assert records[2]["data"]["Vendor Name"] == "Vendor Beta Steel"
    assert records[2]["data"]["Gross Amount"] == 94400.0
    assert records[2]["data"]["TDS (Rs.)"] == 800.0

def test_parse_file_delegation_xlsx(tmp_path):
    """O. Verify parse_file() delegation correctly routes 'xlsx' file_type."""
    p = tmp_path / "delegation.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Test"
    ws.append(["Col1"])
    ws.append(["Val1"])
    wb.save(str(p))
    
    records = parse_file(str(p), "xlsx")
    assert len(records) == 1
    assert records[0]["sheet_name"] == "Test"
    assert records[0]["data"]["Col1"] == "Val1"
