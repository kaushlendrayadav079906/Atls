"""Unit and integration tests for Stock Configuration API and Service."""

import io
import pytest
import openpyxl
from fastapi.testclient import TestClient

from app.main import app
from app.services.stock_configuration_service import StockConfigurationService

client = TestClient(app)


@pytest.fixture
def auth_headers():
    from app.core.security import create_access_token
    token = create_access_token({"sub": "admin", "role": "admin", "branch_id": "01", "permissions": ["products.view"]})
    return {"Authorization": f"Bearer {token}"}


def test_get_stock_configuration_kpis_and_pagination(auth_headers):
    response = client.get("/api/v1/stock-configuration?page=1&page_size=10", headers=auth_headers)
    assert response.status_code == 200
    data = response.json()

    assert "items" in data
    assert "kpis" in data
    assert "filters" in data
    assert len(data["items"]) <= 10
    assert data["kpis"]["total_items"] >= 26
    assert data["kpis"]["active_items"] >= 24
    assert data["kpis"]["inactive_items"] >= 2
    assert data["kpis"]["warehouses_count"] >= 5


def test_update_stock_configuration_item(auth_headers):
    update_payload = {
        "qty": 115.0,
        "grpo_no": "5099",
        "rcvd_qty": 115.0,
        "nt_weight": 14.2,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 110.0,
        "rej_qty": 5.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
    }
    response = client.put("/api/v1/stock-configuration/items/FG001", json=update_payload, headers=auth_headers)
    assert response.status_code == 200
    res_data = response.json()
    assert res_data["success"] is True
    assert res_data["item"]["qty"] == 115.0
    assert res_data["item"]["grpo_no"] == "5099"
    assert res_data["item"]["nt_weight"] == 14.2


def test_add_and_delete_stock_item(auth_headers):
    new_item = {
        "item_code": "TEST999",
        "description": "Test Integration Item",
        "qty": 50.0,
        "stock_qty": 45.0,
        "vendor_name": "Test Vendor",
        "grpo_no": "9999",
        "rcvd_qty": 50.0,
        "nt_weight": 2.5,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 45.0,
        "rej_qty": 5.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Test Group",
        "warehouse_code": "01",
        "is_active": True,
    }
    # Add
    res = client.post("/api/v1/stock-configuration/items", json=new_item, headers=auth_headers)
    assert res.status_code == 200
    assert res.json()["item"]["item_code"] == "TEST999"

    # Delete
    del_res = client.delete("/api/v1/stock-configuration/items/TEST999", headers=auth_headers)
    assert del_res.status_code == 200
    assert del_res.json()["success"] is True


def test_export_stock_configuration_excel(auth_headers):
    response = client.get("/api/v1/stock-configuration/export", headers=auth_headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    # Verify openpyxl can load the exported bytes
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws.title == "Stock Items"
    headers = [cell.value for cell in ws[1]]
    assert "SAP CODE" in headers
    assert "Desc." in headers
    assert "VENDOR NAME" in headers
    assert "GRPO No." in headers
    assert "MC or PT" in headers


def test_import_stock_configuration_excel_preview_and_commit(auth_headers):
    # Create sample workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        "SR. NO.", "SAP CODE", "Desc.", "Qty.", "Stock Qty.", "VENDOR NAME",
        "GRPO No.", "RCVD. Qty.", "NT. WEIGHT", "QC", "DRW. Y/N", "V CHK",
        "Ok Qty.", "Rej. Qty.", "LM", "OLNG", "LPK", "LBL", "MC or PT"
    ])
    ws.append([
        1, "IMP001", "Imported Test Widget", 30, 28, "Vendor X",
        "6001", 30, 1.2, "Y", "Y", "Y",
        28, 2, "Y", "Y", "Y", "Y", "Y"
    ])
    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    # 1. Preview mode
    res_preview = client.post(
        "/api/v1/stock-configuration/import",
        files={"file": ("test_import.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"preview": "true"},
        headers=auth_headers,
    )
    assert res_preview.status_code == 200
    preview_data = res_preview.json()
    assert preview_data["preview"] is True
    assert preview_data["valid_count"] == 1
    assert preview_data["invalid_count"] == 0

    # 2. Commit mode
    res_commit = client.post(
        "/api/v1/stock-configuration/import",
        files={"file": ("test_import.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"preview": "false"},
        headers=auth_headers,
    )
    assert res_commit.status_code == 200
    commit_data = res_commit.json()
    assert commit_data["success"] is True
    assert commit_data["imported_count"] == 1

    # Cleanup imported test item
    client.delete("/api/v1/stock-configuration/items/IMP001", headers=auth_headers)
