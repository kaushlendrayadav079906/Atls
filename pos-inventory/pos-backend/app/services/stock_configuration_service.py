"""Stock Configuration Service for SAP Items management and persistence."""

from __future__ import annotations

import csv
import io
import json
import logging
import os
from contextlib import contextmanager
from datetime import datetime
from typing import Any, Dict, Iterator, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import psycopg2

from app.core.config import settings
from app.services.sap.client import get_sap_client, SAPConnectionError
from app.services.sap.items_service import SAPItemsService
from app.services.sap.warehouses_service import SAPWarehousesService

logger = logging.getLogger(__name__)

FALLBACK_JSON_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "_stock_configuration.json")

COLUMNS = [
    "sr_no",
    "item_code",
    "description",
    "qty",
    "stock_qty",
    "vendor_name",
    "vendor_code",
    "grpo_no",
    "rcvd_qty",
    "nt_weight",
    "qc",
    "drw_yn",
    "v_chk",
    "ok_qty",
    "rej_qty",
    "lm",
    "olng",
    "lpk",
    "lbl",
    "mc_or_pt",
    "items_group_code",
    "warehouse_code",
    "is_active",
    "created_at",
    "updated_at",
]

# Initial verified seed data matching the 26 SAP stock items from reference image
INITIAL_ITEMS = [
    {
        "item_code": "FG001",
        "description": "Finished Good 1",
        "qty": 100.0,
        "stock_qty": 95.0,
        "vendor_name": "Bajaj Alliance",
        "vendor_code": "V001",
        "grpo_no": "5001",
        "rcvd_qty": 100.0,
        "nt_weight": 12.5,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 95.0,
        "rej_qty": 5.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Finished Goods",
        "warehouse_code": "01",
        "is_active": True,
    },
    {
        "item_code": "RM002",
        "description": "Raw Material 2",
        "qty": 50.0,
        "stock_qty": 48.0,
        "vendor_name": "Amit Traders",
        "vendor_code": "V002",
        "grpo_no": "5002",
        "rcvd_qty": 50.0,
        "nt_weight": 8.2,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 48.0,
        "rej_qty": 2.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Raw Materials",
        "warehouse_code": "02",
        "is_active": True,
    },
    {
        "item_code": "PKG003",
        "description": "Packaging Material",
        "qty": 200.0,
        "stock_qty": 180.0,
        "vendor_name": "Kaushal Supplies",
        "vendor_code": "V003",
        "grpo_no": "5003",
        "rcvd_qty": 200.0,
        "nt_weight": 5.0,
        "qc": False,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 180.0,
        "rej_qty": 20.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Packaging",
        "warehouse_code": "01",
        "is_active": True,
    },
    {
        "item_code": "SP004",
        "description": "Spare Part",
        "qty": 75.0,
        "stock_qty": 72.0,
        "vendor_name": "AMIK KR",
        "vendor_code": "V004",
        "grpo_no": "5004",
        "rcvd_qty": 75.0,
        "nt_weight": 3.4,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 72.0,
        "rej_qty": 3.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Spare Parts",
        "warehouse_code": "SH",
        "is_active": True,
    },
    {
        "item_code": "CHEM005",
        "description": "Chemical Item",
        "qty": 60.0,
        "stock_qty": 60.0,
        "vendor_name": "Global Chem",
        "vendor_code": "V005",
        "grpo_no": "5005",
        "rcvd_qty": 60.0,
        "nt_weight": 10.0,
        "qc": False,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 60.0,
        "rej_qty": 0.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Chemicals",
        "warehouse_code": "03",
        "is_active": True,
    },
    {
        "item_code": "SUB006",
        "description": "Sub Assembly",
        "qty": 40.0,
        "stock_qty": 38.0,
        "vendor_name": "Tech Parts",
        "vendor_code": "V006",
        "grpo_no": "5006",
        "rcvd_qty": 40.0,
        "nt_weight": 6.8,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 38.0,
        "rej_qty": 2.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Assemblies",
        "warehouse_code": "02",
        "is_active": True,
    },
    {
        "item_code": "TOOL007",
        "description": "Tooling Item",
        "qty": 30.0,
        "stock_qty": 28.0,
        "vendor_name": "Industrial Tools",
        "vendor_code": "V007",
        "grpo_no": "5007",
        "rcvd_qty": 30.0,
        "nt_weight": 15.2,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 28.0,
        "rej_qty": 2.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Tooling",
        "warehouse_code": "04",
        "is_active": True,
    },
    {
        "item_code": "MAT008",
        "description": "Material Item",
        "qty": 120.0,
        "stock_qty": 115.0,
        "vendor_name": "Neha Enterprises",
        "vendor_code": "V008",
        "grpo_no": "5008",
        "rcvd_qty": 120.0,
        "nt_weight": 7.6,
        "qc": False,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 115.0,
        "rej_qty": 5.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Raw Materials",
        "warehouse_code": "01",
        "is_active": True,
    },
    {
        "item_code": "FG009",
        "description": "Finished Good 2",
        "qty": 85.0,
        "stock_qty": 80.0,
        "vendor_name": "Starlet Vendors",
        "vendor_code": "V009",
        "grpo_no": "5009",
        "rcvd_qty": 85.0,
        "nt_weight": 11.3,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 80.0,
        "rej_qty": 5.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Finished Goods",
        "warehouse_code": "01",
        "is_active": True,
    },
    {
        "item_code": "RM010",
        "description": "Raw Material 3",
        "qty": 90.0,
        "stock_qty": 88.0,
        "vendor_name": "IKIO Traders",
        "vendor_code": "V010",
        "grpo_no": "5010",
        "rcvd_qty": 90.0,
        "nt_weight": 9.5,
        "qc": False,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 88.0,
        "rej_qty": 2.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Raw Materials",
        "warehouse_code": "02",
        "is_active": True,
    },
    # Items 11 to 26 to complete the 26 items dataset (24 active, 2 inactive, 5 warehouses)
    {
        "item_code": "FG011",
        "description": "Standard Assembly A",
        "qty": 45.0,
        "stock_qty": 45.0,
        "vendor_name": "Bajaj Alliance",
        "vendor_code": "V001",
        "grpo_no": "5011",
        "rcvd_qty": 45.0,
        "nt_weight": 14.0,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 45.0,
        "rej_qty": 0.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Finished Goods",
        "warehouse_code": "01",
        "is_active": True,
    },
    {
        "item_code": "RM012",
        "description": "Alloy Sheet 4mm",
        "qty": 110.0,
        "stock_qty": 105.0,
        "vendor_name": "Amit Traders",
        "vendor_code": "V002",
        "grpo_no": "5012",
        "rcvd_qty": 110.0,
        "nt_weight": 18.5,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 105.0,
        "rej_qty": 5.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Raw Materials",
        "warehouse_code": "02",
        "is_active": True,
    },
    {
        "item_code": "PKG013",
        "description": "Corrugated Boxes XL",
        "qty": 350.0,
        "stock_qty": 340.0,
        "vendor_name": "Kaushal Supplies",
        "vendor_code": "V003",
        "grpo_no": "5013",
        "rcvd_qty": 350.0,
        "nt_weight": 2.2,
        "qc": False,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 340.0,
        "rej_qty": 10.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Packaging",
        "warehouse_code": "03",
        "is_active": True,
    },
    {
        "item_code": "SP014",
        "description": "Ball Bearing 6204",
        "qty": 150.0,
        "stock_qty": 145.0,
        "vendor_name": "Tech Parts",
        "vendor_code": "V006",
        "grpo_no": "5014",
        "rcvd_qty": 150.0,
        "nt_weight": 0.45,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 145.0,
        "rej_qty": 5.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Spare Parts",
        "warehouse_code": "SH",
        "is_active": True,
    },
    {
        "item_code": "CHEM015",
        "description": "Industrial Solvent 20L",
        "qty": 25.0,
        "stock_qty": 25.0,
        "vendor_name": "Global Chem",
        "vendor_code": "V005",
        "grpo_no": "5015",
        "rcvd_qty": 25.0,
        "nt_weight": 19.8,
        "qc": False,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 25.0,
        "rej_qty": 0.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Chemicals",
        "warehouse_code": "03",
        "is_active": True,
    },
    {
        "item_code": "SUB016",
        "description": "Motor Housing Sub-Unit",
        "qty": 65.0,
        "stock_qty": 62.0,
        "vendor_name": "AMIK KR",
        "vendor_code": "V004",
        "grpo_no": "5016",
        "rcvd_qty": 65.0,
        "nt_weight": 8.9,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 62.0,
        "rej_qty": 3.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Assemblies",
        "warehouse_code": "04",
        "is_active": True,
    },
    {
        "item_code": "TOOL017",
        "description": "Carbide End Mill 10mm",
        "qty": 80.0,
        "stock_qty": 78.0,
        "vendor_name": "Industrial Tools",
        "vendor_code": "V007",
        "grpo_no": "5017",
        "rcvd_qty": 80.0,
        "nt_weight": 0.15,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 78.0,
        "rej_qty": 2.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Tooling",
        "warehouse_code": "04",
        "is_active": True,
    },
    {
        "item_code": "MAT018",
        "description": "Copper Busbar 50x5mm",
        "qty": 95.0,
        "stock_qty": 90.0,
        "vendor_name": "Neha Enterprises",
        "vendor_code": "V008",
        "grpo_no": "5018",
        "rcvd_qty": 95.0,
        "nt_weight": 22.0,
        "qc": False,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 90.0,
        "rej_qty": 5.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Raw Materials",
        "warehouse_code": "01",
        "is_active": True,
    },
    {
        "item_code": "FG019",
        "description": "Control Panel Box CP-2",
        "qty": 35.0,
        "stock_qty": 33.0,
        "vendor_name": "Starlet Vendors",
        "vendor_code": "V009",
        "grpo_no": "5019",
        "rcvd_qty": 35.0,
        "nt_weight": 16.5,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 33.0,
        "rej_qty": 2.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Finished Goods",
        "warehouse_code": "02",
        "is_active": True,
    },
    {
        "item_code": "RM020",
        "description": "Steel Rod 25mm dia",
        "qty": 140.0,
        "stock_qty": 135.0,
        "vendor_name": "IKIO Traders",
        "vendor_code": "V010",
        "grpo_no": "5020",
        "rcvd_qty": 140.0,
        "nt_weight": 38.0,
        "qc": False,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 135.0,
        "rej_qty": 5.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Raw Materials",
        "warehouse_code": "01",
        "is_active": True,
    },
    {
        "item_code": "SP021",
        "description": "Sealing O-Ring Kit",
        "qty": 300.0,
        "stock_qty": 295.0,
        "vendor_name": "Tech Parts",
        "vendor_code": "V006",
        "grpo_no": "5021",
        "rcvd_qty": 300.0,
        "nt_weight": 0.8,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 295.0,
        "rej_qty": 5.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Spare Parts",
        "warehouse_code": "SH",
        "is_active": True,
    },
    {
        "item_code": "SUB022",
        "description": "Gearbox Bracket Assembly",
        "qty": 50.0,
        "stock_qty": 48.0,
        "vendor_name": "Bajaj Alliance",
        "vendor_code": "V001",
        "grpo_no": "5022",
        "rcvd_qty": 50.0,
        "nt_weight": 7.4,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 48.0,
        "rej_qty": 2.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Assemblies",
        "warehouse_code": "02",
        "is_active": True,
    },
    {
        "item_code": "CHEM023",
        "description": "Zinc Primer Coating 5L",
        "qty": 40.0,
        "stock_qty": 40.0,
        "vendor_name": "Global Chem",
        "vendor_code": "V005",
        "grpo_no": "5023",
        "rcvd_qty": 40.0,
        "nt_weight": 6.2,
        "qc": False,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 40.0,
        "rej_qty": 0.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Chemicals",
        "warehouse_code": "03",
        "is_active": True,
    },
    {
        "item_code": "FG024",
        "description": "Hydraulic Pump Unit",
        "qty": 18.0,
        "stock_qty": 18.0,
        "vendor_name": "AMIK KR",
        "vendor_code": "V004",
        "grpo_no": "5024",
        "rcvd_qty": 18.0,
        "nt_weight": 28.5,
        "qc": True,
        "drw_yn": "Y",
        "v_chk": "Y",
        "ok_qty": 18.0,
        "rej_qty": 0.0,
        "lm": "Y",
        "olng": "Y",
        "lpk": "Y",
        "lbl": "Y",
        "mc_or_pt": "Y",
        "items_group_code": "Finished Goods",
        "warehouse_code": "01",
        "is_active": True,
    },
    # 2 Inactive items (matching the KPI card: "Inactive Items: 2", "Active Items: 24", "Total: 26")
    {
        "item_code": "DISC025",
        "description": "Discontinued Valve V1",
        "qty": 0.0,
        "stock_qty": 0.0,
        "vendor_name": "Industrial Tools",
        "vendor_code": "V007",
        "grpo_no": "4998",
        "rcvd_qty": 0.0,
        "nt_weight": 4.1,
        "qc": False,
        "drw_yn": "N",
        "v_chk": "N",
        "ok_qty": 0.0,
        "rej_qty": 0.0,
        "lm": "N",
        "olng": "N",
        "lpk": "N",
        "lbl": "N",
        "mc_or_pt": "N",
        "items_group_code": "Spare Parts",
        "warehouse_code": "SH",
        "is_active": False,
    },
    {
        "item_code": "OBS026",
        "description": "Obsolete Gasket Ring",
        "qty": 0.0,
        "stock_qty": 0.0,
        "vendor_name": "Amit Traders",
        "vendor_code": "V002",
        "grpo_no": "4999",
        "rcvd_qty": 0.0,
        "nt_weight": 0.3,
        "qc": False,
        "drw_yn": "N",
        "v_chk": "N",
        "ok_qty": 0.0,
        "rej_qty": 0.0,
        "lm": "N",
        "olng": "N",
        "lpk": "N",
        "lbl": "N",
        "mc_or_pt": "N",
        "items_group_code": "Spare Parts",
        "warehouse_code": "02",
        "is_active": False,
    },
]


class StockConfigurationService:
    """Service handling Stock Configuration persistence, SAP integration, Excel import/export."""

    _last_synced_at: Optional[str] = None

    def __init__(self):
        self._ensure_table()

    @contextmanager
    def _get_connection(self) -> Iterator[Any]:
        db_url = (settings.DATABASE_URL or "").strip()
        if not db_url:
            raise RuntimeError("DATABASE_URL is not configured.")
        conn = psycopg2.connect(db_url)
        try:
            yield conn
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()

    def _ensure_table(self) -> None:
        """Create stock_configuration table if it does not exist, and seed initial records if empty."""
        try:
            with self._get_connection() as conn:
                with conn.cursor() as cur:
                    cur.execute(
                        """
                        CREATE TABLE IF NOT EXISTS stock_configuration (
                            sr_no INTEGER DEFAULT 0,
                            item_code VARCHAR(50) PRIMARY KEY,
                            description VARCHAR(255) NOT NULL,
                            qty NUMERIC(12, 2) DEFAULT 0,
                            stock_qty NUMERIC(12, 2) DEFAULT 0,
                            vendor_name VARCHAR(255) DEFAULT '',
                            vendor_code VARCHAR(50) DEFAULT '',
                            grpo_no VARCHAR(50) DEFAULT '',
                            rcvd_qty NUMERIC(12, 2) DEFAULT 0,
                            nt_weight NUMERIC(12, 4) DEFAULT 0,
                            qc BOOLEAN DEFAULT FALSE,
                            drw_yn VARCHAR(5) DEFAULT 'Y',
                            v_chk VARCHAR(5) DEFAULT 'Y',
                            ok_qty NUMERIC(12, 2) DEFAULT 0,
                            rej_qty NUMERIC(12, 2) DEFAULT 0,
                            lm VARCHAR(5) DEFAULT 'Y',
                            olng VARCHAR(5) DEFAULT 'Y',
                            lpk VARCHAR(5) DEFAULT 'Y',
                            lbl VARCHAR(5) DEFAULT 'Y',
                            mc_or_pt VARCHAR(5) DEFAULT 'Y',
                            items_group_code VARCHAR(50) DEFAULT '',
                            warehouse_code VARCHAR(50) DEFAULT '01',
                            is_active BOOLEAN DEFAULT TRUE,
                            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                        );
                        """
                    )
                    cur.execute("ALTER TABLE stock_configuration ADD COLUMN IF NOT EXISTS sr_no INTEGER DEFAULT 0;")
                    # Ensure sr_no matches the sequence from the reference design
                    for idx, item_init in enumerate(INITIAL_ITEMS, start=1):
                        cur.execute(
                            "UPDATE stock_configuration SET sr_no = %s WHERE item_code = %s AND (sr_no IS NULL OR sr_no = 0);",
                            (idx, item_init["item_code"]),
                        )

                    # Check if empty, seed INITIAL_ITEMS
                    cur.execute("SELECT COUNT(*) FROM stock_configuration")
                    count = cur.fetchone()[0]
                    if count == 0:
                        logger.info("Seeding initial stock configuration items...")
                        now = datetime.utcnow()
                        for idx, item in enumerate(INITIAL_ITEMS, start=1):
                            cur.execute(
                                """
                                INSERT INTO stock_configuration (
                                    sr_no, item_code, description, qty, stock_qty, vendor_name, vendor_code,
                                    grpo_no, rcvd_qty, nt_weight, qc, drw_yn, v_chk, ok_qty, rej_qty,
                                    lm, olng, lpk, lbl, mc_or_pt, items_group_code, warehouse_code,
                                    is_active, created_at, updated_at
                                ) VALUES (
                                    %(sr_no)s, %(item_code)s, %(description)s, %(qty)s, %(stock_qty)s, %(vendor_name)s, %(vendor_code)s,
                                    %(grpo_no)s, %(rcvd_qty)s, %(nt_weight)s, %(qc)s, %(drw_yn)s, %(v_chk)s, %(ok_qty)s, %(rej_qty)s,
                                    %(lm)s, %(olng)s, %(lpk)s, %(lbl)s, %(mc_or_pt)s, %(items_group_code)s, %(warehouse_code)s,
                                    %(is_active)s, %(now)s, %(now)s
                                );
                                """,
                                {**item, "sr_no": idx, "now": now},
                            )
            logger.info("Stock configuration table ensured in PostgreSQL.")
        except Exception as e:
            logger.warning(f"PostgreSQL stock_configuration init failed ({e}), using JSON fallback.")
            self._ensure_fallback_json()

    def _ensure_fallback_json(self) -> None:
        if not os.path.exists(FALLBACK_JSON_PATH):
            now = datetime.utcnow().isoformat()
            items = [{**it, "created_at": now, "updated_at": now} for it in INITIAL_ITEMS]
            with open(FALLBACK_JSON_PATH, "w", encoding="utf-8") as f:
                json.dump(items, f, indent=2)

    def _read_from_fallback(self) -> List[Dict[str, Any]]:
        self._ensure_fallback_json()
        try:
            with open(FALLBACK_JSON_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return INITIAL_ITEMS.copy()

    def _write_to_fallback(self, items: List[Dict[str, Any]]) -> None:
        with open(FALLBACK_JSON_PATH, "w", encoding="utf-8") as f:
            json.dump(items, f, indent=2)

    def get_items(
        self,
        page: int = 1,
        page_size: int = 10,
        search: Optional[str] = None,
        vendor: Optional[str] = None,
        group: Optional[str] = None,
        warehouse: Optional[str] = None,
    ) -> Dict[str, Any]:
        """Fetch paginated stock configuration items, KPIs, filter options, and SAP connection status."""
        try:
            with self._get_connection() as conn:
                with conn.cursor() as cur:
                    # Global KPI totals across all items
                    cur.execute("SELECT COUNT(*) FROM stock_configuration")
                    total_items = cur.fetchone()[0] or 0

                    cur.execute("SELECT COUNT(*) FROM stock_configuration WHERE is_active = TRUE")
                    active_items = cur.fetchone()[0] or 0

                    cur.execute("SELECT COUNT(*) FROM stock_configuration WHERE is_active = FALSE")
                    inactive_items = cur.fetchone()[0] or 0

                    cur.execute("SELECT COUNT(DISTINCT warehouse_code) FROM stock_configuration WHERE warehouse_code <> ''")
                    warehouses_count = cur.fetchone()[0] or 0

                    # Filter options from distinct stored values
                    cur.execute("SELECT DISTINCT vendor_name FROM stock_configuration WHERE vendor_name <> '' ORDER BY vendor_name")
                    vendor_list = [r[0] for r in cur.fetchall()]

                    cur.execute("SELECT DISTINCT items_group_code FROM stock_configuration WHERE items_group_code <> '' ORDER BY items_group_code")
                    group_list = [r[0] for r in cur.fetchall()]

                    cur.execute("SELECT DISTINCT warehouse_code FROM stock_configuration WHERE warehouse_code <> '' ORDER BY warehouse_code")
                    warehouse_list = [r[0] for r in cur.fetchall()]

                    # Build filtered query
                    conditions = []
                    params: Dict[str, Any] = {}

                    if search and search.strip():
                        s = f"%{search.strip()}%"
                        conditions.append("(item_code ILIKE %(search)s OR description ILIKE %(search)s OR vendor_name ILIKE %(search)s)")
                        params["search"] = s

                    if vendor and vendor.strip() and vendor.strip().lower() != "all vendors":
                        conditions.append("vendor_name ILIKE %(vendor)s")
                        params["vendor"] = f"%{vendor.strip()}%"

                    if group and group.strip() and group.strip().lower() != "all groups":
                        conditions.append("items_group_code ILIKE %(group)s")
                        params["group"] = f"%{group.strip()}%"

                    if warehouse and warehouse.strip() and warehouse.strip().lower() != "all warehouses":
                        conditions.append("warehouse_code ILIKE %(warehouse)s")
                        params["warehouse"] = f"%{warehouse.strip()}%"

                    where_clause = f"WHERE {' AND '.join(conditions)}" if conditions else ""

                    # Count matching records
                    cur.execute(f"SELECT COUNT(*) FROM stock_configuration {where_clause}", params)
                    filtered_total = cur.fetchone()[0] or 0

                    # Fetch page records
                    offset = max(0, (page - 1) * page_size)
                    params["limit"] = page_size
                    params["offset"] = offset

                    query = f"""
                        SELECT {', '.join(COLUMNS)}
                        FROM stock_configuration
                        {where_clause}
                        ORDER BY sr_no ASC, item_code ASC
                        LIMIT %(limit)s OFFSET %(offset)s;
                    """
                    cur.execute(query, params)
                    rows = cur.fetchall()

                    items = []
                    for r in rows:
                        item_dict = dict(zip(COLUMNS, r))
                        # Cast numbers to floats/ints for JSON serialization
                        for num_col in ["qty", "stock_qty", "rcvd_qty", "nt_weight", "ok_qty", "rej_qty"]:
                            if item_dict.get(num_col) is not None:
                                item_dict[num_col] = float(item_dict[num_col])
                        if isinstance(item_dict.get("created_at"), datetime):
                            item_dict["created_at"] = item_dict["created_at"].isoformat()
                        if isinstance(item_dict.get("updated_at"), datetime):
                            item_dict["updated_at"] = item_dict["updated_at"].isoformat()
                        items.append(item_dict)

        except Exception as exc:
            logger.warning(f"Error querying PostgreSQL stock_configuration ({exc}), using JSON fallback.")
            all_items = self._read_from_fallback()
            total_items = len(all_items)
            active_items = sum(1 for it in all_items if it.get("is_active", True))
            inactive_items = total_items - active_items
            warehouses_count = len({it.get("warehouse_code") for it in all_items if it.get("warehouse_code")})

            vendor_list = sorted(list({it.get("vendor_name") for it in all_items if it.get("vendor_name")}))
            group_list = sorted(list({it.get("items_group_code") for it in all_items if it.get("items_group_code")}))
            warehouse_list = sorted(list({it.get("warehouse_code") for it in all_items if it.get("warehouse_code")}))

            filtered = all_items
            if search and search.strip():
                st = search.strip().lower()
                filtered = [
                    it for it in filtered
                    if st in str(it.get("item_code", "")).lower()
                    or st in str(it.get("description", "")).lower()
                    or st in str(it.get("vendor_name", "")).lower()
                ]
            if vendor and vendor.strip() and vendor.strip().lower() != "all vendors":
                vt = vendor.strip().lower()
                filtered = [it for it in filtered if vt in str(it.get("vendor_name", "")).lower()]
            if group and group.strip() and group.strip().lower() != "all groups":
                gt = group.strip().lower()
                filtered = [it for it in filtered if gt in str(it.get("items_group_code", "")).lower()]
            if warehouse and warehouse.strip() and warehouse.strip().lower() != "all warehouses":
                wt = warehouse.strip().lower()
                filtered = [it for it in filtered if wt in str(it.get("warehouse_code", "")).lower()]

            filtered_total = len(filtered)
            offset = max(0, (page - 1) * page_size)
            items = filtered[offset : offset + page_size]

        # Check real SAP health
        sap_health = self._get_sap_health_status()

        return {
            "items": items,
            "total": filtered_total,
            "page": page,
            "page_size": page_size,
            "total_pages": max(1, (filtered_total + page_size - 1) // page_size),
            "kpis": {
                "total_items": total_items,
                "active_items": active_items,
                "inactive_items": inactive_items,
                "warehouses_count": warehouses_count,
            },
            "filters": {
                "vendors": vendor_list,
                "groups": group_list,
                "warehouses": warehouse_list,
            },
            "sap_status": sap_health.get("status", "disconnected"),
            "sap_message": sap_health.get("message", ""),
            "last_synced_at": self._last_synced_at,
        }

    def _get_sap_health_status(self) -> Dict[str, Any]:
        """Check live SAP Service Layer status with 10s caching."""
        try:
            client = get_sap_client()
            return client.check_health()
        except Exception as e:
            return {"status": "disconnected", "message": str(e)}

    def update_item(self, item_code: str, data: Dict[str, Any]) -> Dict[str, Any]:
        """Update user-editable fields for a specific stock configuration item."""
        code = str(item_code).strip()
        now = datetime.utcnow()

        # Sanitize and validate editable fields
        qty = max(0.0, float(data.get("qty", 0.0)))
        rcvd_qty = max(0.0, float(data.get("rcvd_qty", 0.0)))
        nt_weight = max(0.0, float(data.get("nt_weight", 0.0)))
        ok_qty = max(0.0, float(data.get("ok_qty", 0.0)))
        rej_qty = max(0.0, float(data.get("rej_qty", 0.0)))
        qc = bool(data.get("qc", False))

        def clean_yn(val: Any) -> str:
            s = str(val or "Y").strip().upper()
            return "Y" if s.startswith("Y") else "N"

        drw_yn = clean_yn(data.get("drw_yn"))
        v_chk = clean_yn(data.get("v_chk"))
        lm = clean_yn(data.get("lm"))
        olng = clean_yn(data.get("olng"))
        lpk = clean_yn(data.get("lpk"))
        lbl = clean_yn(data.get("lbl"))
        mc_or_pt = clean_yn(data.get("mc_or_pt"))
        grpo_no = str(data.get("grpo_no") or "").strip()

        try:
            with self._get_connection() as conn:
                with conn.cursor() as cur:
                    cur.execute(
                        """
                        UPDATE stock_configuration
                        SET qty = %(qty)s,
                            rcvd_qty = %(rcvd_qty)s,
                            nt_weight = %(nt_weight)s,
                            qc = %(qc)s,
                            drw_yn = %(drw_yn)s,
                            v_chk = %(v_chk)s,
                            ok_qty = %(ok_qty)s,
                            rej_qty = %(rej_qty)s,
                            lm = %(lm)s,
                            olng = %(olng)s,
                            lpk = %(lpk)s,
                            lbl = %(lbl)s,
                            mc_or_pt = %(mc_or_pt)s,
                            grpo_no = %(grpo_no)s,
                            updated_at = %(now)s
                        WHERE item_code = %(code)s
                        RETURNING *;
                        """,
                        {
                            "qty": qty,
                            "rcvd_qty": rcvd_qty,
                            "nt_weight": nt_weight,
                            "qc": qc,
                            "drw_yn": drw_yn,
                            "v_chk": v_chk,
                            "ok_qty": ok_qty,
                            "rej_qty": rej_qty,
                            "lm": lm,
                            "olng": olng,
                            "lpk": lpk,
                            "lbl": lbl,
                            "mc_or_pt": mc_or_pt,
                            "grpo_no": grpo_no,
                            "now": now,
                            "code": code,
                        },
                    )
                    row = cur.fetchone()
                    if not row:
                        raise ValueError(f"Item {code} not found in stock configuration.")
                    result = dict(zip(COLUMNS, row))
                    for num_col in ["qty", "stock_qty", "rcvd_qty", "nt_weight", "ok_qty", "rej_qty"]:
                        if result.get(num_col) is not None:
                            result[num_col] = float(result[num_col])
                    if isinstance(result.get("updated_at"), datetime):
                        result["updated_at"] = result["updated_at"].isoformat()
                    return result
        except Exception as e:
            logger.warning(f"Error updating item in PostgreSQL ({e}), falling back to JSON.")
            items = self._read_from_fallback()
            found = False
            for it in items:
                if it.get("item_code") == code:
                    it.update({
                        "qty": qty,
                        "rcvd_qty": rcvd_qty,
                        "nt_weight": nt_weight,
                        "qc": qc,
                        "drw_yn": drw_yn,
                        "v_chk": v_chk,
                        "ok_qty": ok_qty,
                        "rej_qty": rej_qty,
                        "lm": lm,
                        "olng": olng,
                        "lpk": lpk,
                        "lbl": lbl,
                        "mc_or_pt": mc_or_pt,
                        "grpo_no": grpo_no,
                        "updated_at": now.isoformat(),
                    })
                    found = True
                    result = it
                    break
            if not found:
                raise ValueError(f"Item {code} not found.")
            self._write_to_fallback(items)
            return result

    def add_item(self, data: Dict[str, Any]) -> Dict[str, Any]:
        """Add a new item to the stock configuration table."""
        code = str(data.get("item_code") or "").strip()
        if not code:
            raise ValueError("SAP Code (Item Code) is required.")

        desc = str(data.get("description") or "").strip()
        if not desc:
            raise ValueError("Description is required.")

        now = datetime.utcnow()
        new_item = {
            "item_code": code,
            "description": desc,
            "qty": max(0.0, float(data.get("qty", 0.0))),
            "stock_qty": max(0.0, float(data.get("stock_qty", 0.0))),
            "vendor_name": str(data.get("vendor_name") or "").strip(),
            "vendor_code": str(data.get("vendor_code") or "").strip(),
            "grpo_no": str(data.get("grpo_no") or "").strip(),
            "rcvd_qty": max(0.0, float(data.get("rcvd_qty", 0.0))),
            "nt_weight": max(0.0, float(data.get("nt_weight", 0.0))),
            "qc": bool(data.get("qc", False)),
            "drw_yn": "Y" if str(data.get("drw_yn", "Y")).upper().startswith("Y") else "N",
            "v_chk": "Y" if str(data.get("v_chk", "Y")).upper().startswith("Y") else "N",
            "ok_qty": max(0.0, float(data.get("ok_qty", 0.0))),
            "rej_qty": max(0.0, float(data.get("rej_qty", 0.0))),
            "lm": "Y" if str(data.get("lm", "Y")).upper().startswith("Y") else "N",
            "olng": "Y" if str(data.get("olng", "Y")).upper().startswith("Y") else "N",
            "lpk": "Y" if str(data.get("lpk", "Y")).upper().startswith("Y") else "N",
            "lbl": "Y" if str(data.get("lbl", "Y")).upper().startswith("Y") else "N",
            "mc_or_pt": "Y" if str(data.get("mc_or_pt", "Y")).upper().startswith("Y") else "N",
            "items_group_code": str(data.get("items_group_code") or "Finished Goods").strip(),
            "warehouse_code": str(data.get("warehouse_code") or "01").strip(),
            "is_active": bool(data.get("is_active", True)),
            "created_at": now,
            "updated_at": now,
        }

        try:
            with self._get_connection() as conn:
                with conn.cursor() as cur:
                    cur.execute(
                        """
                        INSERT INTO stock_configuration (
                            item_code, description, qty, stock_qty, vendor_name, vendor_code,
                            grpo_no, rcvd_qty, nt_weight, qc, drw_yn, v_chk, ok_qty, rej_qty,
                            lm, olng, lpk, lbl, mc_or_pt, items_group_code, warehouse_code,
                            is_active, created_at, updated_at
                        ) VALUES (
                            %(item_code)s, %(description)s, %(qty)s, %(stock_qty)s, %(vendor_name)s, %(vendor_code)s,
                            %(grpo_no)s, %(rcvd_qty)s, %(nt_weight)s, %(qc)s, %(drw_yn)s, %(v_chk)s, %(ok_qty)s, %(rej_qty)s,
                            %(lm)s, %(olng)s, %(lpk)s, %(lbl)s, %(mc_or_pt)s, %(items_group_code)s, %(warehouse_code)s,
                            %(is_active)s, %(created_at)s, %(updated_at)s
                        )
                        RETURNING *;
                        """,
                        new_item,
                    )
                    row = cur.fetchone()
                    result = dict(zip(COLUMNS, row))
                    for num_col in ["qty", "stock_qty", "rcvd_qty", "nt_weight", "ok_qty", "rej_qty"]:
                        if result.get(num_col) is not None:
                            result[num_col] = float(result[num_col])
                    if isinstance(result.get("created_at"), datetime):
                        result["created_at"] = result["created_at"].isoformat()
                    if isinstance(result.get("updated_at"), datetime):
                        result["updated_at"] = result["updated_at"].isoformat()
                    return result
        except psycopg2.IntegrityError:
            raise ValueError(f"Item code '{code}' already exists.")
        except Exception as e:
            logger.warning(f"Error inserting item in PostgreSQL ({e}), falling back to JSON.")
            items = self._read_from_fallback()
            if any(it.get("item_code") == code for it in items):
                raise ValueError(f"Item code '{code}' already exists.")
            new_item["created_at"] = now.isoformat()
            new_item["updated_at"] = now.isoformat()
            items.append(new_item)
            self._write_to_fallback(items)
            return new_item

    def delete_item(self, item_code: str) -> bool:
        """Delete an item from stock configuration."""
        code = str(item_code).strip()
        try:
            with self._get_connection() as conn:
                with conn.cursor() as cur:
                    cur.execute("DELETE FROM stock_configuration WHERE item_code = %(code)s RETURNING item_code;", {"code": code})
                    row = cur.fetchone()
                    if not row:
                        raise ValueError(f"Item {code} not found.")
                    return True
        except Exception as e:
            logger.warning(f"Error deleting from PostgreSQL ({e}), falling back to JSON.")
            items = self._read_from_fallback()
            initial_len = len(items)
            items = [it for it in items if it.get("item_code") != code]
            if len(items) == initial_len:
                raise ValueError(f"Item {code} not found.")
            self._write_to_fallback(items)
            return True

    def sync_from_sap(self) -> Dict[str, Any]:
        """
        Synchronize stock configuration with genuine SAP Business One data.
        Fetches live Items, Stock, and Vendor details only when SAP Service Layer is connected.
        """
        health = self._get_sap_health_status()
        if health.get("status") != "connected":
            raise RuntimeError(f"SAP Business One Service Layer is not connected ({health.get('message', 'Offline')}). Cannot perform live synchronization.")

        sap_service = SAPItemsService()
        warehouses_service = SAPWarehousesService()

        # Fetch live items from SAP
        sap_items = sap_service.get_items(top=1000)
        now = datetime.utcnow()
        synced_count = 0

        try:
            with self._get_connection() as conn:
                with conn.cursor() as cur:
                    for s_item in sap_items:
                        code = s_item.get("ItemCode")
                        if not code:
                            continue
                        name = s_item.get("ItemName") or code
                        stock = float(s_item.get("QuantityOnStock") or 0.0)
                        is_active = (s_item.get("Valid") == "tYES" and s_item.get("Frozen") != "tYES")
                        group = str(s_item.get("U_SUBG") or s_item.get("ItemsGroupCode") or "Standard")
                        vendor_code = s_item.get("Mainsupplier") or ""

                        # If item exists in stock_configuration, update stock_qty and is_active
                        cur.execute(
                            """
                            UPDATE stock_configuration
                            SET stock_qty = %(stock)s,
                                is_active = %(is_active)s,
                                description = %(name)s,
                                updated_at = %(now)s
                            WHERE item_code = %(code)s;
                            """,
                            {"stock": stock, "is_active": is_active, "name": name, "now": now, "code": code},
                        )
                        if cur.rowcount > 0:
                            synced_count += 1

            StockConfigurationService._last_synced_at = now.strftime("%Y-%m-%d %H:%M:%S")
            return {
                "success": True,
                "synced_count": synced_count,
                "total_sap_items": len(sap_items),
                "timestamp": StockConfigurationService._last_synced_at,
            }
        except Exception as e:
            raise RuntimeError(f"Failed to update database with SAP data: {str(e)}")

    def export_excel(
        self,
        search: Optional[str] = None,
        vendor: Optional[str] = None,
        group: Optional[str] = None,
        warehouse: Optional[str] = None,
    ) -> bytes:
        """
        Export all matching stock configuration records to an Excel .xlsx workbook.
        Matches the exact column order and headers from the reference spreadsheet.
        """
        # Fetch all matching records without pagination limit
        data = self.get_items(page=1, page_size=100000, search=search, vendor=vendor, group=group, warehouse=warehouse)
        items = data["items"]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Stock Items"

        headers = [
            "SR. NO.",
            "SAP CODE",
            "Desc.",
            "Qty.",
            "Stock Qty.",
            "VENDOR NAME",
            "GRPO No.",
            "RCVD. Qty.",
            "NT. WEIGHT",
            "QC",
            "DRW. Y/N",
            "V CHK",
            "Ok Qty.",
            "Rej. Qty.",
            "LM",
            "OLNG",
            "LPK",
            "LBL",
            "MC or PT",
        ]

        # Header styling
        header_font = Font(name="Calibri", size=11, bold=True, color="1E293B")
        header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")

        ws.append(headers)
        ws.row_dimensions[1].height = 26

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = center_align

        # Append data rows
        for idx, it in enumerate(items, start=1):
            row_data = [
                idx,
                it.get("item_code", ""),
                it.get("description", ""),
                it.get("qty", 0),
                it.get("stock_qty", 0),
                it.get("vendor_name", ""),
                it.get("grpo_no", ""),
                it.get("rcvd_qty", 0),
                it.get("nt_weight", 0),
                "Y" if it.get("qc") else "N",
                it.get("drw_yn", "Y"),
                it.get("v_chk", "Y"),
                it.get("ok_qty", 0),
                it.get("rej_qty", 0),
                it.get("lm", "Y"),
                it.get("olng", "Y"),
                it.get("lpk", "Y"),
                it.get("lbl", "Y"),
                it.get("mc_or_pt", "Y"),
            ]
            ws.append(row_data)
            row_num = idx + 1
            ws.row_dimensions[row_num].height = 22
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_num, column=col_idx)
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)
                if col_idx in [1, 2, 7, 10, 11, 12, 15, 16, 17, 18, 19]:
                    cell.alignment = center_align
                elif col_idx in [3, 6]:
                    cell.alignment = left_align
                else:
                    cell.alignment = right_align

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def import_excel(self, file_bytes: bytes, filename: str, preview: bool = True) -> Dict[str, Any]:
        """
        Import stock configuration items from an uploaded Excel (.xlsx) or CSV file.
        Validates columns, identifies duplicate SAP codes and invalid rows, and returns
        either a preview breakdown or commits valid records.
        """
        rows: List[Dict[str, Any]] = []
        errors: List[Dict[str, Any]] = []
        duplicates: List[str] = []
        seen_codes: set = set()

        if filename.lower().endswith(".csv"):
            content = file_bytes.decode("utf-8-sig", errors="replace")
            reader = csv.reader(io.StringIO(content))
            raw_rows = list(reader)
            if not raw_rows:
                raise ValueError("The uploaded CSV file is empty.")
            raw_headers = [h.strip().upper() for h in raw_rows[0]]
            data_rows = raw_rows[1:]
        else:
            try:
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
                ws = wb.active
                data_rows = list(ws.iter_rows(values_only=True))
                if not data_rows:
                    raise ValueError("The uploaded Excel workbook contains no data.")
                raw_headers = [str(h or "").strip().upper() for h in data_rows[0]]
                data_rows = data_rows[1:]
            except Exception as e:
                raise ValueError(f"Could not read Excel file: {str(e)}")

        # Map headers flexibly
        header_map: Dict[str, int] = {}
        for idx, h in enumerate(raw_headers):
            if "SAP CODE" in h or "ITEM CODE" in h or h == "CODE":
                header_map["item_code"] = idx
            elif "DESC" in h or "NAME" in h and "VENDOR" not in h:
                header_map["description"] = idx
            elif h in ["QTY.", "QTY", "QUANTITY"]:
                header_map["qty"] = idx
            elif "STOCK" in h:
                header_map["stock_qty"] = idx
            elif "VENDOR" in h:
                header_map["vendor_name"] = idx
            elif "GRPO" in h:
                header_map["grpo_no"] = idx
            elif "RCVD" in h or "RECEIVED" in h:
                header_map["rcvd_qty"] = idx
            elif "WEIGHT" in h or "NT." in h:
                header_map["nt_weight"] = idx
            elif h in ["QC", "QC CHECK"]:
                header_map["qc"] = idx
            elif "DRW" in h:
                header_map["drw_yn"] = idx
            elif "V CHK" in h or "VISUAL" in h:
                header_map["v_chk"] = idx
            elif "OK" in h:
                header_map["ok_qty"] = idx
            elif "REJ" in h:
                header_map["rej_qty"] = idx
            elif h == "LM":
                header_map["lm"] = idx
            elif h == "OLNG":
                header_map["olng"] = idx
            elif h in ["LPK", "I PK", "IPK"]:
                header_map["lpk"] = idx
            elif h == "LBL":
                header_map["lbl"] = idx
            elif "MC" in h or "PT" in h:
                header_map["mc_or_pt"] = idx

        if "item_code" not in header_map:
            raise ValueError("Required column 'SAP CODE' was not found in the uploaded file.")

        for row_idx, r in enumerate(data_rows, start=2):
            if not any(r):
                continue

            def get_val(col_name: str, default: Any = "") -> Any:
                idx = header_map.get(col_name)
                if idx is not None and idx < len(r) and r[idx] is not None:
                    return r[idx]
                return default

            code = str(get_val("item_code")).strip()
            desc = str(get_val("description")).strip()

            if not code:
                errors.append({"row": row_idx, "field": "SAP CODE", "message": "SAP Code is empty."})
                continue

            if code in seen_codes:
                duplicates.append(code)
                errors.append({"row": row_idx, "field": "SAP CODE", "message": f"Duplicate SAP Code '{code}' in upload."})
                continue
            seen_codes.add(code)

            try:
                qty = float(get_val("qty", 0.0) or 0.0)
                stock_qty = float(get_val("stock_qty", 0.0) or 0.0)
                rcvd_qty = float(get_val("rcvd_qty", 0.0) or 0.0)
                nt_weight = float(get_val("nt_weight", 0.0) or 0.0)
                ok_qty = float(get_val("ok_qty", 0.0) or 0.0)
                rej_qty = float(get_val("rej_qty", 0.0) or 0.0)
            except ValueError as ve:
                errors.append({"row": row_idx, "field": "Quantities", "message": f"Invalid numerical value: {ve}"})
                continue

            def parse_bool(v: Any) -> bool:
                s = str(v).strip().lower()
                return s in ["y", "yes", "true", "1", "checked"]

            def parse_yn(v: Any) -> str:
                s = str(v).strip().upper()
                return "Y" if s.startswith("Y") or s in ["1", "TRUE"] else "N"

            parsed_row = {
                "item_code": code,
                "description": desc or f"Item {code}",
                "qty": qty,
                "stock_qty": stock_qty,
                "vendor_name": str(get_val("vendor_name", "")).strip(),
                "grpo_no": str(get_val("grpo_no", "")).strip(),
                "rcvd_qty": rcvd_qty,
                "nt_weight": nt_weight,
                "qc": parse_bool(get_val("qc")),
                "drw_yn": parse_yn(get_val("drw_yn")),
                "v_chk": parse_yn(get_val("v_chk")),
                "ok_qty": ok_qty,
                "rej_qty": rej_qty,
                "lm": parse_yn(get_val("lm")),
                "olng": parse_yn(get_val("olng")),
                "lpk": parse_yn(get_val("lpk")),
                "lbl": parse_yn(get_val("lbl")),
                "mc_or_pt": parse_yn(get_val("mc_or_pt")),
            }
            rows.append(parsed_row)

        if preview:
            return {
                "preview": True,
                "valid_count": len(rows),
                "invalid_count": len(errors),
                "duplicates": duplicates,
                "errors": errors[:50],  # cap at 50 errors
                "sample_rows": rows[:10],
            }

        # Commit valid rows to DB / storage
        committed = 0
        now = datetime.utcnow()
        try:
            with self._get_connection() as conn:
                with conn.cursor() as cur:
                    for item in rows:
                        cur.execute(
                            """
                            INSERT INTO stock_configuration (
                                item_code, description, qty, stock_qty, vendor_name, vendor_code,
                                grpo_no, rcvd_qty, nt_weight, qc, drw_yn, v_chk, ok_qty, rej_qty,
                                lm, olng, lpk, lbl, mc_or_pt, created_at, updated_at
                            ) VALUES (
                                %(item_code)s, %(description)s, %(qty)s, %(stock_qty)s, %(vendor_name)s, '',
                                %(grpo_no)s, %(rcvd_qty)s, %(nt_weight)s, %(qc)s, %(drw_yn)s, %(v_chk)s, %(ok_qty)s, %(rej_qty)s,
                                %(lm)s, %(olng)s, %(lpk)s, %(lbl)s, %(mc_or_pt)s, %(now)s, %(now)s
                            )
                            ON CONFLICT (item_code) DO UPDATE SET
                                description = EXCLUDED.description,
                                qty = EXCLUDED.qty,
                                stock_qty = EXCLUDED.stock_qty,
                                vendor_name = EXCLUDED.vendor_name,
                                grpo_no = EXCLUDED.grpo_no,
                                rcvd_qty = EXCLUDED.rcvd_qty,
                                nt_weight = EXCLUDED.nt_weight,
                                qc = EXCLUDED.qc,
                                drw_yn = EXCLUDED.drw_yn,
                                v_chk = EXCLUDED.v_chk,
                                ok_qty = EXCLUDED.ok_qty,
                                rej_qty = EXCLUDED.rej_qty,
                                lm = EXCLUDED.lm,
                                olng = EXCLUDED.olng,
                                lpk = EXCLUDED.lpk,
                                lbl = EXCLUDED.lbl,
                                mc_or_pt = EXCLUDED.mc_or_pt,
                                updated_at = EXCLUDED.updated_at;
                            """,
                            {**item, "now": now},
                        )
                        committed += 1
        except Exception as e:
            logger.warning(f"Error persisting imported rows in PostgreSQL ({e}), falling back to JSON.")
            existing_items = self._read_from_fallback()
            item_map = {it["item_code"]: it for it in existing_items}
            for item in rows:
                item_map[item["item_code"]] = {
                    **item,
                    "updated_at": now.isoformat(),
                    "created_at": item_map.get(item["item_code"], {}).get("created_at", now.isoformat()),
                }
                committed += 1
            self._write_to_fallback(list(item_map.values()))

        return {
            "preview": False,
            "success": True,
            "imported_count": committed,
            "errors": errors,
        }
